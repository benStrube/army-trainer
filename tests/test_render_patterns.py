from pathlib import Path

import pytest
from pptx import Presentation

from army_trainer.plan.spec import SlideSpec
from army_trainer.render.deck import render_slides
from army_trainer.render.patterns import REGISTRY
from army_trainer.render.template import build_template

FIXTURE = Path(__file__).parent / "fixtures" / "spec_all_patterns.json"
FM = Path(__file__).parents[1] / "specs" / "FM-3-09.spec.json"


@pytest.fixture(scope="module")
def tpl(tmp_path_factory):
    return build_template(tmp_path_factory.mktemp("t") / "base.pptx")


def test_registry_covers_every_pattern():
    spec = SlideSpec.model_validate_json(FIXTURE.read_text())
    assert {s.pattern for s in spec.slides} <= set(REGISTRY)


def test_fixture_renders_every_pattern_without_overflow(tmp_path, tpl):
    spec = SlideSpec.model_validate_json(FIXTURE.read_text())
    out = tmp_path / "d.pptx"
    ctx = render_slides(spec, out, tpl)
    assert len(Presentation(out).slides) == len(spec.slides)
    assert ctx.warnings == []


def test_fm_3_09_spec_renders(tmp_path, tpl):
    """Every slide draws; fit warnings (est. overflow at the font floor) are for 3.3 to split."""
    spec = SlideSpec.model_validate_json(FM.read_text())
    out = tmp_path / "d.pptx"
    render_slides(spec, out, tpl)
    assert len(Presentation(out).slides) == len(spec.slides)


def test_text_is_copied_exactly(tmp_path, tpl):
    spec = SlideSpec.model_validate_json(FIXTURE.read_text())
    out = tmp_path / "d.pptx"
    render_slides(spec, out, tpl)
    texts = {
        t_
        for sl in Presentation(out).slides
        for sh in sl.shapes
        if sh.has_text_frame
        for t_ in [sh.text_frame.text]
    }
    for s in spec.slides:
        if s.pattern in {"takeaways", "checklist", "whats_new"}:
            for it in s.items:
                assert it.text in texts
