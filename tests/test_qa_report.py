"""Local review report (WP 4.3)."""

import json

import pytest
from test_plan_packet_check import DOC, GLOSSARY, checklist, deck, item
from test_qa_rules import PAD

from army_trainer.index.build import build_indexes
from army_trainer.plan.spec import SlideSpec
from army_trainer.qa.report import build_report, deck_slide_map
from army_trainer.qa.review import Review, claims, sha256_file
from army_trainer.qa.rules import run_qa
from army_trainer.structure.parse import parse_markdown


@pytest.fixture(scope="module")
def tree():
    return parse_markdown(DOC + GLOSSARY)


@pytest.fixture
def spec_path(tmp_path):
    slide = checklist(
        item("The FO will locate <b>targets</b>.", ["para-1-1"], "will"),  # source says must
        item("Check the message type and location.", ["para-1-2.li1"]),
        *PAD[2:],
    )
    p = tmp_path / "FM-0-00.spec.json"
    p.write_text(json.dumps(deck([slide])))
    return p


def _review(spec_path, **kw):
    ids = [c.id for c in claims(SlideSpec.model_validate_json(spec_path.read_text()))]
    base = {
        "pub_id": "FM-0-00",
        "source_sha256": "abc",
        "spec_sha256": sha256_file(spec_path),
        "reviewed_on": "2026-10-02",
        "reviewer": "test",
        "verdicts": {i: "pass" for i in ids},
    }
    return Review.model_validate(base | kw)


def test_report_shows_claims_source_flags_and_verdicts(tree, spec_path, tmp_path):
    qa = run_qa(spec_path, tree, build_indexes(tree))
    rp = tmp_path / "FM-0-00.review.json"
    rp.write_text(_review(spec_path).model_dump_json())
    page = build_report(spec_path, tree, qa, tmp_path / "rep", None, rp).read_text()
    assert 'id="s02"' in page and "The FO is the fire support representative" in page
    assert "[directive]" in page and "is not in the cited text" in page  # rule flag on the claim
    assert 'class="badge pass"' in page  # review verdict
    assert "&lt;b&gt;targets&lt;/b&gt;" in page and "<b>targets</b>" not in page  # escaped
    assert "changed after this review" not in page
    assert "no thumbnails" in page


def test_report_warns_when_the_review_is_stale(tree, spec_path, tmp_path):
    qa = run_qa(spec_path, tree, build_indexes(tree))
    rp = tmp_path / "FM-0-00.review.json"
    rp.write_text(_review(spec_path).model_dump_json())
    spec_path.write_text(spec_path.read_text() + " ")
    page = build_report(spec_path, tree, qa, tmp_path / "rep", None, rp).read_text()
    assert "changed after this review" in page


def test_report_without_review_or_deck(tree, spec_path, tmp_path):
    qa = run_qa(spec_path, tree, build_indexes(tree))
    page = build_report(spec_path, tree, qa, tmp_path / "rep").read_text()
    assert "No fidelity review file" in page


def test_deck_slide_map_follows_split_slides(tmp_path):
    from pptx import Presentation

    spec = SlideSpec.model_validate(deck([checklist(*PAD)]))
    titles = ["Test Manual", "Checks (1 of 2)", "Checks (2 of 2)", "Read it"]
    prs = Presentation()
    for t in titles:
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = t
    p = tmp_path / "d.pptx"
    prs.save(p)
    assert deck_slide_map(spec, p) == {"s01": [1], "s02": [2, 3], "s03": [4]}
    assert deck_slide_map(spec, None) == {}
