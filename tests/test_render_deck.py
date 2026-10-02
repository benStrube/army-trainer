from pathlib import Path

import pytest
from pptx import Presentation

from army_trainer.plan.spec import SlideSpec
from army_trainer.render.deck import _pieces, footer_text, render_deck
from army_trainer.render.template import build_template
from army_trainer.structure.models import DocTree

ROOT = Path(__file__).parents[1]
TREE = ROOT / "data/json/FM-3-09.json"
SPEC = ROOT / "specs/FM-3-09.spec.json"
needs_tree = pytest.mark.skipif(not TREE.exists(), reason="run convert for FM-3-09 first")


def test_footer_text_groups_and_truncates():
    assert footer_text("FM 3-09", ["para 2-19"], "T") == "FM 3-09 · para 2-19 · T"
    cites = [f"para 1-{i}" for i in range(1, 9)] + ["table 4-1", "p. vii", "p. ix"]
    out = footer_text("FM 3-09", cites, "T")
    assert out == "FM 3-09 · paras 1-1, 1-2, 1-3, 1-4, 1-5 +3 more · table 4-1, p. vii +1 more · T"
    assert footer_text("FM 3-09", [], "T") == "FM 3-09 · T"


def _spec():
    return SlideSpec.model_validate_json(SPEC.read_text())


def test_overflowing_list_slides_split_until_they_fit():
    s = next(x for x in _spec().slides if x.pattern == "acronyms")
    big = s.model_copy(update={"entries": s.entries * 3})  # 3x the entries
    parts = _pieces(big, lambda p: len(p.entries) <= 6)
    assert len(parts) > 1
    assert [e for p in parts for e in p.entries] == big.entries  # nothing lost or reordered
    assert all(len(p.entries) >= 4 for p in parts)


def test_unsplittable_pattern_is_never_split():
    s = next(x for x in _spec().slides if x.pattern == "process_flow")
    assert _pieces(s, lambda p: False) == [s]


@needs_tree
def test_full_deck_has_footers_notes_and_verbatim_source(tmp_path):
    tree = DocTree.model_validate_json(TREE.read_text())
    tpl = build_template(tmp_path / "base.pptx")
    out = tmp_path / "deck.pptx"
    res = render_deck(_spec(), tree, out, tpl)
    prs = Presentation(out)
    assert len(prs.slides) == res.slide_count >= len(_spec().slides)
    footer = next(p for p in prs.slides[1].placeholders if p.placeholder_format.type == 15)
    assert footer.text_frame.text.startswith("FM 3-09 · ")
    assert footer.text_frame.text.endswith("Unofficial training aid")
    # slide 3 (takeaways): notes quote para 2-19 verbatim from the tree
    note = prs.slides[2].notes_slide.notes_text_frame.text
    para = next(n for d in tree.divisions for n in _walk(d) if getattr(n, "number", None) == "2-19")
    assert para.plain in note
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                assert "Army star" not in sh.text_frame.text


def _walk(node):
    yield node
    for c in getattr(node, "children", None) or []:
        yield from _walk(c)
