from pptx import Presentation

from army_trainer.render import theme as t
from army_trainer.render.template import build_template


def test_contrast_pairs_meet_threshold():
    for label, (fg, bg, minimum) in t.TEXT_PAIRS.items():
        ratio = t.contrast(getattr(t.PALETTE, fg), getattr(t.PALETTE, bg))
        assert ratio >= minimum, f"{label}: {ratio:.2f}"


def test_gold_text_on_white_is_the_known_bad_pair():
    assert t.contrast(t.PALETTE.army_gold, t.PALETTE.white) < 2


def test_template_layouts_and_fonts(tmp_path):
    path = build_template(tmp_path / "base.pptx")
    prs = Presentation(path)
    assert [la.name for la in prs.slide_layouts] == [
        "Title Slide",
        "Section Header",
        "Content",
        "Blank",
    ]
    assert abs(prs.slide_width / 914400 - t.SLIDE_W) < 0.01
    content = next(la for la in prs.slide_layouts if la.name == "Content")
    assert content.placeholders[0].placeholder_format.idx == 0
    slide = prs.slides.add_slide(content)
    slide.shapes.title.text = "Hello"
    xml = prs.slide_master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
    ).blob.decode()
    assert t.PALETTE.army_gold in xml and "Arial" in xml
