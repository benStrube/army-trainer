"""Rule-based QA checks (WP 4.2): `army-trainer qa <ID>`.

Deterministic, no API calls. Six checks run on the committed spec (plus the rendered deck when
it exists). They flag things for a person or the Opus fidelity review to look at; they don't
replace that review (D12).

* citation: every item cites nodes that exist, cites are paragraphs not whole headings, content
  slides carry cites, and the rendered deck shows the claim text, a citation footer and notes.
* verbatim: numbers, dates and form numbers on a slide appear in the cited text.
* directive: directive words (will / must / will not / may ...) on a slide appear in the cited
  text; an item's `directive` is in the source; the rendered deck keeps the directive word.
* readability (D15): per-slide Flesch-Kincaid grade with glossary terms and acronyms counted as
  one word; the deck passes when >= 90% of scored slides are at grade 9.0 or below.
* acronym: acronyms are in the glossary, spelled out where first used, and listed on the
  acronyms slide.
* coverage: share of the publication's will / must requirements that a slide cites, per chapter,
  and the ones no slide cites.

Findings reuse the plan check's rules (`plan.check`) so the two never disagree.
"""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from ..plan.check import _cited_objects, _slide_text, _strings, check_spec
from ..plan.nodes import NodeIndex
from ..plan.spec import STRUCTURAL_PATTERNS, SlideSpec
from ..structure.models import DocTree
from .review import claims, deck_text, rendered

CHECKS = ("citation", "verbatim", "directive", "readability", "acronym", "coverage")
#: Flesch-Kincaid grade above which one statement is flagged (target ~8th grade).
GRADE_WARN = 12.0
#: A slide (mean of its statements) is "at level" at or below this grade (D15).
GRADE_SLIDE = 9.0
#: The deck passes when at least this share of scored slides are at level (MVP criterion 8.4).
SLIDE_SHARE_PASS = 0.90
#: Slides whose text is fixed, a list of expansions or pointers elsewhere (D15).
READABILITY_EXEMPT = frozenset({"title", "acronyms", "closing"})
MIN_WORDS = 8  # shorter text has no stable grade level
#: Share of mandatory requirements a chapter's slides should cite before coverage warns.
COVERAGE_WARN = 0.10

_FORM_RX = re.compile(
    r"\b(?:(?:DA|DD|SF)\s+Form|DA\s+Pam|FM|AR|ATP|ADP|TC|TM|JP)\s+\d[\w-]*(?:\.\d+)?", re.I
)
_MON = (
    r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?"
    r"|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
)
# Months are capitalized; "May" (also a directive word) only counts with a day or year.
_DATE_RX = re.compile(
    rf"\b(?:\d{{1,2}}\s+)?{_MON}\b\.?(?:\s+\d{{1,2}}(?!\d))?(?:,?\s+\d{{4}})?"
    r"|\b(?:\d{1,2}\s+)?May(?:\s+\d{1,2},?\s+\d{4}|\s+\d{4}|\s+\d{1,2}\b)"
    r"|\b\d{1,2}\s+May\b"
    r"|\b\d{1,2}/\d{1,2}/\d{2,4}\b|\b(?:19|20)\d{2}\b"
)


@dataclass
class QaFinding:
    check: str
    level: str  # "error" | "warning"
    where: str
    message: str

    def __str__(self) -> str:
        return f"{self.level.upper():7} [{self.check}] {self.where}: {self.message}"


@dataclass
class QaReport:
    pub_id: str
    findings: list[QaFinding] = field(default_factory=list)
    coverage: dict = field(default_factory=dict)
    readability: dict = field(default_factory=dict)
    deck_checked: bool = False

    @property
    def errors(self) -> list[QaFinding]:
        return [f for f in self.findings if f.level == "error"]

    @property
    def warnings(self) -> list[QaFinding]:
        return [f for f in self.findings if f.level == "warning"]

    def by_check(self, name: str) -> list[QaFinding]:
        return [f for f in self.findings if f.check == name]

    def to_json(self) -> dict:
        return {
            "pub_id": self.pub_id,
            "deck_checked": self.deck_checked,
            "errors": len(self.errors),
            "warnings": len(self.warnings),
            "readability": self.readability,
            "coverage": self.coverage,
            "findings": [asdict(f) for f in self.findings],
        }


# ---------------------------------------------------------------- helpers


def _norm_form(s: str) -> str:
    return re.sub(r"\s+", " ", s.upper()).strip()


def forms_and_dates(text: str) -> tuple[set[str], set[str]]:
    """(form / publication numbers, dates) in `text`, normalized for comparison."""
    forms = {_norm_form(m) for m in _FORM_RX.findall(text)}
    text = _FORM_RX.sub(" ", text)  # "DA Form 2028" is not the year 2028
    dates = {re.sub(r"[\s.,]+", " ", m).strip().lower() for m in _DATE_RX.findall(text)}
    return forms, dates


def _map_plan_finding(f) -> QaFinding:
    """Put a `plan.check` finding under the QA check it belongs to."""
    m = f.message
    if "not in the cited text" in m and m.startswith("number"):
        check = "verbatim"
    elif m.startswith(("directive", "slide text says", "title says")):
        check = "directive"
    elif "acronym" in m:
        check = "acronym"
    elif m.startswith("term ") or "meaning differs" in m:
        check = "acronym"
    else:
        check = "citation"
    return QaFinding(check, f.level, f.where, m)


# ---------------------------------------------------------------- checks


def check_forms_dates(spec: SlideSpec, idx: NodeIndex) -> list[QaFinding]:
    """Form numbers, publication numbers and dates on a slide must be in the cited text."""
    out = []
    for s in spec.slides:
        for path, obj in _cited_objects(s, s.id):
            known = [c for c in obj.cite if c in idx]
            if not known:
                continue
            source = " ".join(idx.text(c) for c in known)
            sf, sd = forms_and_dates(source)
            tf, td = forms_and_dates(_slide_text(obj))
            where = path.replace(f"{s.id}.", f"{s.id} ", 1) if path != s.id else s.id
            for x in sorted(tf - sf):
                out.append(
                    QaFinding(
                        "verbatim", "error", where, f"form/pub number {x!r} not in the cited text"
                    )
                )
            for x in sorted(td - sd):
                out.append(
                    QaFinding("verbatim", "error", where, f"date {x!r} not in the cited text")
                )
    return out


def check_uncited(spec: SlideSpec) -> list[QaFinding]:
    """Content slides must carry at least one cited item (structural slides are exempt)."""
    out = []
    for s in spec.slides:
        if s.pattern in STRUCTURAL_PATTERNS:
            continue
        cited = list(_cited_objects(s, s.id))
        if not cited:
            out.append(QaFinding("citation", "error", s.id, "content slide has no cited items"))
        for path, obj in cited:
            if not obj.cite:
                out.append(QaFinding("citation", "error", path, "item has an empty cite list"))
    return out


def neutralizer(idx: NodeIndex | None):
    """Regex replacing the publication's glossary terms and acronyms with one word (D15)."""
    phrases, abbrs = set(), set()
    for loc in idx.by_id.values() if idx else ():
        n = loc.node
        if n.type == "term" and n.term.strip():
            phrases.add(n.term.strip())
        elif n.type == "acronym":
            if n.meaning.strip():
                phrases.add(n.meaning.strip())
            if n.abbreviation.strip():
                abbrs.add(n.abbreviation.strip())
    if not phrases and not abbrs:
        return None
    # longest first, whole words only; phrases ignore case, abbreviations don't
    ph = "|".join(re.escape(p) for p in sorted(phrases, key=len, reverse=True))
    ab = "|".join(re.escape(a) for a in sorted(abbrs, key=len, reverse=True))
    rx_p = re.compile(rf"(?<![\w-])(?:{ph})(?![\w-])", re.I) if ph else None
    rx_a = re.compile(rf"(?<![\w-])(?:{ab})(?![\w-])") if ab else None

    def neutralize(text: str) -> str:
        if rx_p:
            text = rx_p.sub("term", text)
        if rx_a:
            text = rx_a.sub("term", text)
        return text

    return neutralize


def check_readability(
    spec: SlideSpec, idx: NodeIndex | None = None
) -> tuple[list[QaFinding], dict]:
    """Per-slide, term-aware Flesch-Kincaid grade (D15, `docs/decisions/readability.md`)."""
    import textstat

    neutralize = (neutralizer(idx) if idx is not None else None) or (lambda t: t)
    out: list[QaFinding] = []
    slides: dict[str, dict] = {}
    statements = []  # neutralized grades
    by_slide: dict[str, list[tuple[float, float]]] = {}
    for c in claims(spec):
        if c.in_notes or c.pattern in READABILITY_EXEMPT:
            continue
        for part in c.parts:
            if len(re.findall(r"[A-Za-z']+", part)) < MIN_WORDS:
                continue
            g = textstat.flesch_kincaid_grade(neutralize(part))
            raw = textstat.flesch_kincaid_grade(part)
            statements.append(g)
            by_slide.setdefault(c.slide, []).append((g, raw))
            if g > GRADE_WARN:
                out.append(
                    QaFinding(
                        "readability", "warning", c.id,
                        f"grade level {g:.1f} (> {GRADE_WARN:.0f}): {part[:80]}",
                    )
                )  # fmt: skip
    for s in spec.slides:
        rows = by_slide.get(s.id)
        if not rows:
            continue
        grade = sum(g for g, _ in rows) / len(rows)
        slides[s.id] = {
            "grade": round(grade, 2),
            "raw_grade": round(sum(r for _, r in rows) / len(rows), 2),
            "statements": len(rows),
        }
        if grade > GRADE_SLIDE:
            out.append(
                QaFinding(
                    "readability", "warning", s.id,
                    f"slide grade level {grade:.1f} (> {GRADE_SLIDE:.1f})",
                )
            )  # fmt: skip
    ok = sum(1 for v in slides.values() if v["grade"] <= GRADE_SLIDE)
    share = ok / len(slides) if slides else 1.0
    avg = sum(statements) / len(statements) if statements else 0.0
    passed = share >= SLIDE_SHARE_PASS
    raw_ok = sum(1 for v in slides.values() if v["raw_grade"] <= GRADE_SLIDE)
    if not passed:
        out.append(
            QaFinding(
                "readability", "warning", "deck",
                f"{ok} of {len(slides)} slides at grade {GRADE_SLIDE:.0f} or below "
                f"({share:.0%}); target {SLIDE_SHARE_PASS:.0%}: FAIL",
            )
        )  # fmt: skip
    stats = {
        "method": "D15",
        "neutralized": neutralize is not None and idx is not None,
        "statements": len(statements),
        "mean_grade": round(avg, 2),
        "over_limit": sum(1 for g in statements if g > GRADE_WARN),
        "scored_slides": len(slides),
        "slides_ok": ok,
        "slide_share": round(share, 3),
        "target_share": SLIDE_SHARE_PASS,
        "passed": passed,
        "raw_slides_ok": raw_ok,
        "slides": slides,
    }
    return out, stats


def check_acronyms(spec: SlideSpec, idx: NodeIndex) -> list[QaFinding]:
    """A glossary acronym should be spelled out on its first slide (or just before it)."""
    meanings = {}
    for loc in idx.by_id.values():
        n = loc.node
        if n.type == "acronym" and len(n.abbreviation) >= 2:
            meanings[n.abbreviation] = n.meaning
    out, seen = [], set()
    text_so_far = ""
    for s in spec.slides:
        if s.pattern == "acronyms":
            continue
        text = " ".join(_strings(s))
        for abbr, meaning in meanings.items():
            if abbr in seen:
                continue
            if re.search(rf"(?<![\w-]){re.escape(abbr)}(?![\w-])(?!\s+\d)", text):
                seen.add(abbr)
                spelled = meaning.lower() in (text_so_far + " " + text).lower()
                if not spelled:
                    out.append(
                        QaFinding(
                            "acronym", "warning", s.id,
                            f"{abbr} is used before it is spelled out ({meaning!r})",
                        )
                    )  # fmt: skip
        text_so_far += " " + text.lower()
    return out


def check_deck(spec: SlideSpec, idx: NodeIndex, pptx: Path) -> list[QaFinding]:
    """The rendered deck keeps the claims, directive words, footers and the disclaimer."""
    from pptx import Presentation

    out = []
    dtext = deck_text(pptx)
    for c in claims(spec):
        if c.in_notes:
            continue
        if not rendered(c, dtext):
            out.append(
                QaFinding("citation", "error", c.id, "claim text is not in the rendered deck")
            )
        elif c.directive and c.directive.lower() not in dtext:
            out.append(
                QaFinding(
                    "directive", "error", c.id, f"directive {c.directive!r} missing from the deck"
                )
            )
    prs = Presentation(str(pptx))
    no_footer = [s.id for s in spec.slides if s.pattern in STRUCTURAL_PATTERNS]
    n_footed = 0
    for slide in prs.slides:
        text = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame).lower()
        n_footed += "unofficial training aid" in text
    if n_footed < len(prs.slides) - len(no_footer) - 4:  # title/closing/dividers carry it elsewhere
        out.append(
            QaFinding(
                "citation", "warning", "deck",
                f"only {n_footed} of {len(prs.slides)} slides show the unofficial-aid footer",
            )
        )  # fmt: skip
    if "unofficial training aid" not in dtext:
        out.append(
            QaFinding("citation", "error", "deck", "the disclaimer is missing from the deck")
        )
    notes_missing = [
        i + 1
        for i, slide in enumerate(prs.slides)
        if slide.has_notes_slide
        and "SOURCE TEXT" not in slide.notes_slide.notes_text_frame.text
        and "TALKING POINTS" in slide.notes_slide.notes_text_frame.text
    ]
    if notes_missing:
        out.append(
            QaFinding(
                "citation",
                "warning",
                "deck",
                f"slides {notes_missing[:5]} have notes without source text",
            )
        )
    return out


def _covered_ids(spec: SlideSpec, idx: NodeIndex) -> set[str]:
    """Nodes a slide cites, with everything underneath them (a paragraph's list items)."""
    ids: set[str] = set()
    for s in spec.slides:
        cited = {c for _, o in _cited_objects(s, s.id) for c in o.cite if c in idx}
        cited |= {c for c in s.notes.extra_sources if c in idx}
        stack = [idx.get(c) for c in cited]
        while stack:
            n = stack.pop()
            ids.add(n.id)
            stack.extend(getattr(n, "children", None) or [])
    return ids


def coverage(
    spec: SlideSpec, tree: DocTree, indexes, idx: NodeIndex
) -> tuple[list[QaFinding], dict]:
    """How many mandatory requirements (will / must / ...) each chapter's slides cite."""
    covered = _covered_ids(spec, idx)
    per: dict[str, dict] = {}
    missing: list[dict] = []
    for d in indexes.directives:
        if d["strength"] != "mandatory" or d.get("possibly_historical") or d["node_id"] not in idx:
            continue
        div = idx.division_of(d["node_id"])
        row = per.setdefault(div, {"title": idx.get(div).title, "requirements": 0, "cited": 0})
        row["requirements"] += 1
        if d["node_id"] in covered:
            row["cited"] += 1
        else:
            missing.append(
                {
                    "node_id": d["node_id"],
                    "division": div,
                    "cite": d["cite"],
                    "sentence": d["sentence"],
                }
            )
    out = []
    for div, row in per.items():
        row["share"] = round(row["cited"] / row["requirements"], 3) if row["requirements"] else 1.0
        kind = idx.get(div).kind
        if kind == "chapter" and row["requirements"] >= 5 and row["share"] < COVERAGE_WARN:
            out.append(
                QaFinding(
                    "coverage", "warning", div,
                    f"{row['cited']} of {row['requirements']} mandatory requirements cited "
                    f"({row['share']:.0%})",
                )
            )  # fmt: skip
    total = sum(r["requirements"] for r in per.values())
    cited = sum(r["cited"] for r in per.values())
    chapters_without = [
        d.id for d in tree.divisions if d.kind == "chapter" and not any(
            s.chapter == d.id for s in spec.slides
        )
    ]  # fmt: skip
    for cid in chapters_without:
        out.append(QaFinding("coverage", "warning", cid, "no slide belongs to this chapter"))
    return out, {
        "mandatory_requirements": total,
        "cited": cited,
        "share": round(cited / total, 3) if total else 1.0,
        "per_division": per,
        "uncited": missing,
    }


# ---------------------------------------------------------------- run


def run_qa(spec_path: Path, tree: DocTree, indexes, deck: Path | None = None) -> QaReport:
    spec = SlideSpec.model_validate_json(spec_path.read_text())
    report = QaReport(pub_id=spec.pub.pub_id)
    idx = NodeIndex(tree)
    plan = check_spec(json.loads(spec_path.read_text()), tree)
    report.findings += [_map_plan_finding(f) for f in plan]
    report.findings += check_uncited(spec)
    report.findings += check_forms_dates(spec, idx)
    rf, report.readability = check_readability(spec, idx)
    report.findings += rf
    report.findings += check_acronyms(spec, idx)
    cf, report.coverage = coverage(spec, tree, indexes, idx)
    report.findings += cf
    if deck and deck.exists():
        report.deck_checked = True
        report.findings += check_deck(spec, idx, deck)
    # the same issue can arrive from two checks
    seen, unique = set(), []
    for f in report.findings:
        key = (f.check, f.where, f.message)
        if key not in seen:
            seen.add(key)
            unique.append(f)
    report.findings = unique
    return report
