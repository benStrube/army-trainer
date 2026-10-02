"""Local review report (WP 4.3): `army-trainer qa <ID> --report`.

One static HTML page per deck (`out/reports/<ID>/index.html`, thumbnails beside it). For every
slide: the rendered thumbnail on the left; on the right each claim next to the full text of the
nodes it cites, the rule-based flags (WP 4.2) and, when `specs/<ID>.review.json` exists, the
fidelity verdict (WP 4.1). A summary at the top shows the totals, coverage per chapter and the
requirements no slide cites. Nothing is sent anywhere: it is a local file for a person.
"""

from __future__ import annotations

import html
import re
import shutil
import subprocess
from collections import defaultdict
from pathlib import Path

from pptx import Presentation

from ..plan.nodes import NodeIndex
from ..plan.spec import SlideSpec
from .review import Review, claims, sha256_file
from .rules import QaReport

_WHERE_RX = re.compile(r"^(s\d+)(?:[ .]|$)")
_SPLIT_RX = re.compile(r" \(\d+ of \d+\)$")


def make_thumbnails(pptx: Path, out_dir: Path) -> list[Path]:
    """PNG per deck slide via LibreOffice + pdftoppm. Returns [] when the tools are missing."""
    if not (shutil.which("soffice") and shutil.which("pdftoppm")):
        return []
    thumbs = out_dir / "thumbs"
    thumbs.mkdir(parents=True, exist_ok=True)
    for old in thumbs.glob("slide-*.png"):
        old.unlink()
    pdf = thumbs / f"{pptx.stem}.pdf"
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(thumbs), str(pptx)],
        check=True,
        capture_output=True,
        timeout=300,
    )
    if not pdf.exists():  # soffice exits 0 without Impress installed
        return []
    subprocess.run(
        ["pdftoppm", "-png", "-r", "50", str(pdf), str(thumbs / "slide")],
        check=True,
        capture_output=True,
        timeout=300,
    )
    pdf.unlink(missing_ok=True)
    return sorted(thumbs.glob("slide-*.png"))


def deck_slide_map(spec: SlideSpec, pptx: Path | None) -> dict[str, list[int]]:
    """Spec slide id -> deck slide numbers (1-based). A slide the renderer split maps to
    several; titles ending "(1 of 2)" are matched back to the spec title."""
    if pptx is None or not pptx.exists():
        return {}
    titles = []
    for s in Presentation(str(pptx)).slides:
        t = s.shapes.title
        titles.append(_SPLIT_RX.sub("", t.text.strip()) if t is not None else "")
    out: dict[str, list[int]] = {}
    pos = 0
    for sl in spec.slides:
        n = [pos]
        while pos + len(n) < len(titles) and titles[pos + len(n)] == sl.title and len(n) < 6:
            n.append(pos + len(n))
        if titles[pos : pos + 1] != [sl.title] and pos < len(titles):
            n = [pos]  # unmatched title (no title placeholder): keep deck order
        out[sl.id] = [i + 1 for i in n]
        pos += len(n)
    return out


def _e(s: object) -> str:
    return html.escape(str(s), quote=True)


CSS = """
:root{--ink:#1a1a1a;--mut:#5f6368;--line:#d9d9d9;--bg:#fff;--gold:#c99700;--bad:#b3261e;
--warn:#8a5a00;--ok:#1b6b34;--card:#f6f6f6}
@media (prefers-color-scheme:dark){:root{--ink:#eee;--mut:#a8abb0;--line:#3a3a3a;--bg:#161616;
--gold:#e0b030;--bad:#ff8a80;--warn:#f0b64a;--ok:#7fd69a;--card:#202020}}
body{font:15px/1.45 Arial,sans-serif;color:var(--ink);background:var(--bg);margin:0 auto;
max-width:1200px;padding:16px}
h1{font-size:22px;border-bottom:3px solid var(--gold);padding-bottom:6px}
h2{font-size:17px;margin:0}
table{border-collapse:collapse;width:100%}td,th{border:1px solid var(--line);padding:4px 8px;
text-align:left;vertical-align:top}
.slide{border:1px solid var(--line);border-radius:6px;margin:18px 0;background:var(--card)}
.slide>header{padding:8px 12px;border-bottom:1px solid var(--line);display:flex;gap:10px;
flex-wrap:wrap;align-items:baseline}
.body{display:grid;grid-template-columns:minmax(220px,360px) 1fr;gap:14px;padding:12px}
@media (max-width:760px){.body{grid-template-columns:1fr}}
.thumbs{position:sticky;top:8px;align-self:start}
.thumbs img{width:100%;border:1px solid var(--line);margin-bottom:6px;
background:#fff}
.claim{border-top:1px solid var(--line);padding:8px 0}.claim:first-child{border-top:0}
.claim .say{font-weight:bold}
.src{margin:4px 0 0 12px;color:var(--mut);font-size:13px;border-left:3px solid var(--gold);
padding-left:8px}
.badge{display:inline-block;border-radius:10px;padding:0 8px;font-size:12px;border:1px solid
currentColor}
.pass{color:var(--ok)}.minor,.warning{color:var(--warn)}.major,.critical,.error{color:var(--bad)}
.mut{color:var(--mut)}.flag{font-size:13px;margin:2px 0}
.banner{border:2px solid var(--bad);padding:8px;margin:10px 0}
details{margin:8px 0}summary{cursor:pointer}
"""


def _badge(kind: str, label: str | None = None) -> str:
    return f'<span class="badge {_e(kind)}">{_e(label or kind)}</span>'


def build_report(
    spec_path: Path,
    tree,
    qa: QaReport,
    out_dir: Path,
    deck: Path | None = None,
    review_path: Path | None = None,
) -> Path:
    spec = SlideSpec.model_validate_json(spec_path.read_text())
    idx = NodeIndex(tree)
    out_dir.mkdir(parents=True, exist_ok=True)
    thumbs = make_thumbnails(deck, out_dir) if deck and deck.exists() else []
    smap = deck_slide_map(spec, deck)

    review: Review | None = None
    stale = False
    if review_path and review_path.exists():
        review = Review.model_validate_json(review_path.read_text())
        stale = review.spec_sha256 != sha256_file(spec_path)

    flags: dict[str, list] = defaultdict(list)  # slide id -> findings
    flags_by_claim: dict[str, list] = defaultdict(list)
    deck_level = []
    for f in qa.findings:
        m = _WHERE_RX.match(f.where)
        if not m:
            deck_level.append(f)
            continue
        flags[m.group(1)].append(f)
        flags_by_claim[f.where.replace(" ", ".", 1)].append(f)

    by_slide = defaultdict(list)
    for c in claims(spec):
        by_slide[c.slide].append(c)

    h = [
        '<!doctype html><html lang="en"><head><meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width,initial-scale=1">',
        f"<title>{_e(spec.pub.short_name)} review report</title><style>{CSS}</style></head><body>",
        f"<h1>{_e(spec.pub.short_name)}: review report (unofficial training aid)</h1>",
        f'<p class="mut">{len(spec.slides)} spec slides · {len(claims(spec))} claims · '
        f"{len(qa.errors)} rule error(s) · {len(qa.warnings)} warning(s)"
        + ("" if thumbs else " · no thumbnails (LibreOffice or pdftoppm not found, or no deck)")
        + "</p>",
    ]
    if review:
        c = {v: sum(1 for x in review.verdicts.values() if x == v) for v in
             ("pass", "minor", "major", "critical")}  # fmt: skip
        h.append(
            f"<p>Fidelity review ({_e(review.reviewer)}, {_e(review.reviewed_on)}): "
            f"{c['pass']} pass, {c['minor']} minor, {c['major']} major, "
            f"{c['critical']} critical.</p>"
        )
        if stale:
            h.append(
                '<div class="banner">The spec changed after this review: verdicts below are '
                "out of date until an Opus session re-reviews it.</div>"
            )
    else:
        h.append('<p class="mut">No fidelity review file found for this spec.</p>')

    # summary: checks, readability, coverage
    h.append(
        "<h2>Rule-based checks</h2><table><tr><th>Check</th><th>Errors</th><th>Warnings</th></tr>"
    )
    from .rules import CHECKS

    for name in CHECKS:
        fs = qa.by_check(name)
        e = sum(f.level == "error" for f in fs)
        h.append(f"<tr><td>{name}</td><td>{e}</td><td>{len(fs) - e}</td></tr>")
    h.append("</table>")
    r, cov = qa.readability, qa.coverage
    h.append(
        f"<p>Readability (D15): {r.get('slides_ok')} of {r.get('scored_slides')} slides at grade 9 "
        f"or below ({r.get('slide_share', 0):.0%}; target {r.get('target_share', 0):.0%}: "
        f"{'PASS' if r.get('passed') else 'FAIL'}). Neutralized statement mean "
        f"{r.get('mean_grade')} over {r.get('statements')} statements "
        f"({r.get('over_limit')} over grade 12). Coverage: {cov.get('cited')} of "
        f"{cov.get('mandatory_requirements')} mandatory requirements cited "
        f"({cov.get('share', 0):.0%}).</p>"
    )
    rs = r.get("slides", {})
    h.append("<details><summary>Readability by slide (neutralized / raw grade)</summary><table>"
             "<tr><th>Slide</th><th>Grade</th><th>Raw</th><th>Statements</th></tr>")  # fmt: skip
    for sid, row in rs.items():
        warn = ' class="flag"' if row["grade"] > 9.0 else ""
        h.append(
            f"<tr{warn}><td>{_e(sid)}</td><td>{row['grade']}</td><td>{row['raw_grade']}</td>"
            f"<td>{row['statements']}</td></tr>"
        )
    h.append("</table></details>")
    h.append("<details><summary>Coverage by chapter</summary><table><tr><th>Division</th>"
             "<th>Title</th><th>Cited</th><th>Requirements</th></tr>")  # fmt: skip
    for div, row in cov.get("per_division", {}).items():
        h.append(
            f"<tr><td>{_e(div)}</td><td>{_e(row['title'])}</td><td>{row['cited']}</td>"
            f"<td>{row['requirements']}</td></tr>"
        )
    h.append("</table></details>")
    uncited = cov.get("uncited", [])
    h.append(f"<details><summary>{len(uncited)} mandatory requirements no slide cites</summary>")
    cur = None
    for u in uncited:
        if u["division"] != cur:
            cur = u["division"]
            h.append(f"<h3>{_e(cur)}</h3>")
        h.append(f'<div class="flag"><b>{_e(u["cite"])}</b> {_e(u["sentence"])}</div>')
    h.append("</details>")
    if deck_level:
        h.append("<h2>Deck-level flags</h2>")
        h += [f'<div class="flag">{_badge(f.level)} [{_e(f.check)}] {_e(f.where)}: '
              f"{_e(f.message)}</div>" for f in deck_level]  # fmt: skip

    h.append("<h2>Slides</h2>")
    for s in spec.slides:
        nums = smap.get(s.id, [])
        sf = flags.get(s.id, [])
        n_err = sum(f.level == "error" for f in sf)
        h.append(f'<section class="slide" id="{_e(s.id)}"><header><h2>{_e(s.id)} · {_e(s.title)}'
                 f'</h2><span class="mut">{_e(s.pattern)}'
                 + (f" · deck slide {', '.join(map(str, nums))}" if nums else "")
                 + "</span>"
                 + (_badge("error", f"{n_err} error(s)") if n_err else "")
                 + (_badge("warning", f"{len(sf) - n_err} warning(s)") if len(sf) > n_err else "")
                 + '</header><div class="body"><div class="thumbs">')  # fmt: skip
        for n in nums:
            if n <= len(thumbs):
                rel = thumbs[n - 1].relative_to(out_dir).as_posix()
                h.append(f'<img src="{_e(rel)}" alt="Deck slide {n}" loading="lazy">')
        h.append("</div><div>")
        for f in (f for f in sf if f.where == s.id):
            h.append(f'<div class="flag">{_badge(f.level)} [{_e(f.check)}] {_e(f.message)}</div>')
        cs = by_slide.get(s.id, [])
        if not cs:
            h.append('<p class="mut">No cited claims on this slide.</p>')
        for c in cs:
            h.append('<div class="claim">')
            v = review.verdicts.get(c.id) if review else None
            h.append(
                f'<div><span class="say">{_e(c.text)}</span> '
                f'<span class="mut">({_e(c.id)}{"; speaker notes" if c.in_notes else ""})</span> '
                + (_badge(v) if v else "")
                + (f" {_badge('mut', 'directive: ' + c.directive)}" if c.directive else "")
                + "</div>"
            )  # fmt: skip
            for f in flags_by_claim.get(c.id, []):
                h.append(
                    f'<div class="flag">{_badge(f.level)} [{_e(f.check)}] {_e(f.message)}</div>'
                )
            if review:
                for rf in review.findings:
                    if rf.claim == c.id:
                        h.append(
                            f'<div class="flag">{_badge(rf.verdict)} review ({_e(rf.status)}): '
                            f"{_e(rf.note)}</div>"
                        )
            for cite in c.cites:
                node = idx.get(cite)
                h.append(f'<div class="src"><b>{_e(node.cite)}</b> {_e(idx.text(cite))}</div>')
            h.append("</div>")
        h.append("</div></div></section>")
    if review and review.deck_notes:
        h.append("<h2>Deck notes from the review</h2><ul>")
        h += [f"<li>{_e(n)}</li>" for n in review.deck_notes]
        h.append("</ul>")
    h.append("</body></html>")
    out = out_dir / "index.html"
    out.write_text("\n".join(h), encoding="utf-8")
    return out
