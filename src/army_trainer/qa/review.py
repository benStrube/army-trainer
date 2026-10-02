"""Session fidelity review (WP 4.1, D10, Opus-owned).

`army-trainer qa <ID> --review-packet` writes `data/packets/<ID>/review.md`: every claim in the
spec (each cited statement, speaker-note talking points included) next to the **full** text of
the nodes it cites, the requirement sentences in that text, and whether the claim's text made
it into the rendered deck. An Opus session reviews it with `qa/prompts/fidelity_review.md` and
writes `specs/<ID>.review.json` (the `Review` model below). `qa <ID> --review-check` confirms
the review covers every claim of the current spec and reports whether the deck passes.

No Claude API calls (D10): the session reads the packet; this code only prepares and checks.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, ConfigDict, Field

from ..plan.check import TEXT_FIELDS, _cited_objects, _numbers, _slide_text, directive_words
from ..plan.nodes import NodeIndex
from ..plan.spec import SlideSpec
from ..structure.models import DocTree

REVIEW_VERSION = "1.0"
RUBRIC_VERSION = "1.0"
Verdict = Literal["pass", "minor", "major", "critical"]
Category = Literal[
    "added_requirement",  # a will/must/... or an obligation the source doesn't state
    "changed_directive",  # softened, strengthened or dropped directive word
    "number_mismatch",  # number, unit, distance, date or form number differs
    "changed_meaning",  # says something different from the source
    "unsupported",  # the cited text doesn't support it (wrong or missing cite)
    "outside_knowledge",  # true or not, it isn't in the publication
    "misleading_omission",  # drops a condition/exception that changes what a Soldier does
    "citation_scope",  # supported, but the cite is too broad or misses a node
    "plain_language",  # jargon, undefined acronym, too long or too hard
    "render",  # the deck doesn't show the claim as the spec says
]
Status = Literal["fixed", "open", "accepted"]


class Finding(BaseModel):
    model_config = ConfigDict(extra="forbid")

    claim: str = Field(
        description="Claim id from the review packet, e.g. 's09.roles[0].duties[2]'."
    )
    verdict: Literal["minor", "major", "critical"]
    category: Category
    note: str = Field(min_length=1, description="What is wrong, quoting the source.")
    fix: str | None = Field(default=None, description="The change made or proposed.")
    status: Status = "open"


class Review(BaseModel):
    """`specs/<ID>.review.json`."""

    model_config = ConfigDict(extra="forbid")

    review_version: Literal["1.0"] = REVIEW_VERSION
    rubric_version: Literal["1.0"] = RUBRIC_VERSION
    pub_id: str
    source_sha256: str
    spec_sha256: str = Field(
        description="SHA-256 of the spec file reviewed; a later edit voids it."
    )
    reviewed_on: str
    reviewer: str = Field(description="Model tier and mode, e.g. 'Opus session (D10)'.")
    verdicts: dict[str, Verdict] = Field(
        description="Final verdict for every claim id in the packet (after fixes)."
    )
    findings: list[Finding] = Field(
        default_factory=list, description="Every non-pass finding, including fixed ones."
    )
    deck_notes: list[str] = Field(
        default_factory=list, description="Deck-level observations (balance, coverage, render)."
    )


# ---------------------------------------------------------------- claims


@dataclass
class Claim:
    id: str
    slide: str
    pattern: str
    title: str
    text: str
    directive: str | None
    cites: list[str]
    in_notes: bool
    parts: list[str] = field(default_factory=list)  # each text field, as drawn separately
    numbers: list[str] = field(default_factory=list)
    words: list[str] = field(default_factory=list)  # directive words in the claim text


def claims(spec: SlideSpec) -> list[Claim]:
    out = []
    for s in spec.slides:
        for path, obj in _cited_objects(s, s.id):
            text = _slide_text(obj)
            out.append(
                Claim(
                    id=path,
                    slide=s.id,
                    pattern=s.pattern,
                    title=s.title,
                    text=text,
                    directive=getattr(obj, "directive", None),
                    cites=list(obj.cite),
                    in_notes=".notes." in path,
                    parts=_parts(obj),
                    numbers=sorted(_numbers(text)),
                    words=list(dict.fromkeys(directive_words(text))),
                )
            )
    return out


def _parts(obj) -> list[str]:
    out = []
    for name in TEXT_FIELDS:
        v = getattr(obj, name, None)
        out += [v] if isinstance(v, str) else [x for x in (v or []) if isinstance(x, str)]
    return out


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---------------------------------------------------------------- deck text


def _norm(s: str) -> str:
    s = s.replace("’", "'").replace("“", '"').replace("”", '"').replace("•", " ")
    return re.sub(r"\s+", " ", s).strip().lower()


def deck_text(pptx: Path) -> str:
    """All visible text of a rendered deck (not the notes), normalized."""
    from pptx import Presentation

    parts = []
    for slide in Presentation(str(pptx)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                parts.extend(c.text for r in shape.table.rows for c in r.cells)
    return _norm(" ".join(parts))


def rendered(claim: Claim, deck: str) -> bool:
    """Every text field of the claim appears in the deck's visible text."""
    return all(_norm(p) in deck for p in claim.parts)


# ---------------------------------------------------------------- packet


def build_review_packet(
    spec_path: Path, tree: DocTree, indexes, out_dir: Path, deck: Path | None = None
) -> Path:
    spec = SlideSpec.model_validate_json(spec_path.read_text())
    idx = NodeIndex(tree)
    cs = claims(spec)
    dtext = deck_text(deck) if deck and deck.exists() else None
    reqs: dict[str, list[dict]] = {}
    for d in indexes.directives:
        reqs.setdefault(d["node_id"], []).append(d)

    def covered(node_id: str) -> list[str]:
        """The node and every node under it (requirements in a cited paragraph's list items)."""
        out, stack = [], [idx.get(node_id)]
        while stack:
            n = stack.pop()
            out.append(n.id)
            stack.extend(getattr(n, "children", None) or [])
        return out

    out_dir.mkdir(parents=True, exist_ok=True)
    md = [
        f"# Fidelity review packet: {spec.pub.short_name}\n",
        f"- Spec: `{spec_path}` (sha256 `{sha256_file(spec_path)}`)",
        f"- Source sha256: `{spec.pub.source_sha256}`",
        f"- Claims: {len(cs)} on {len(spec.slides)} slides"
        + (f"; deck checked: `{deck}`" if dtext is not None else "; deck not rendered"),
        "",
        "Review with `src/army_trainer/qa/prompts/fidelity_review.md`. Claim ids are the keys "
        "of `verdicts` in the review file.\n",
    ]
    data = []
    current = None
    for c in cs:
        if c.slide != current:
            current = c.slide
            md.append(f"\n## {c.slide} · {c.pattern} · {c.title}\n")
        flags = []
        if c.directive:
            flags.append(f"directive **{c.directive}**")
        if c.words and c.words != [c.directive]:
            flags.append("words: " + ", ".join(c.words))
        if c.numbers:
            flags.append("numbers: " + ", ".join(c.numbers))
        if c.in_notes:
            flags.append("speaker notes")
        is_rendered = None
        if dtext is not None and not c.in_notes:
            is_rendered = rendered(c, dtext)
            if not is_rendered:
                flags.append("**NOT FOUND IN DECK**")
        md.append(f"### `{c.id}`" + (f"  ({'; '.join(flags)})" if flags else ""))
        md.append(f"> {c.text}\n")
        rows = []
        for cite in c.cites:
            node = idx.get(cite)
            md.append(f"- **[{cite}]** ({node.cite}) {idx.text(cite)}")
            for nid in covered(cite):
                rows += [r for r in reqs.get(nid, []) if r["strength"] != "permissive"]
        if rows:
            md.append("- Requirement sentences in the cited text:")
            md.extend(f"  - {r['verb']}: {r['sentence']}" for r in rows[:8])
            if len(rows) > 8:
                md.append(f"  - … {len(rows) - 8} more")
        md.append("")
        data.append(asdict(c) | {"rendered": is_rendered})
    out = out_dir / "review.md"
    out.write_text("\n".join(md) + "\n", encoding="utf-8")
    (out_dir / "review_packet.json").write_text(json.dumps(data, indent=1), encoding="utf-8")
    return out


# ---------------------------------------------------------------- review check


@dataclass
class ReviewReport:
    errors: list[str]
    open_blocking: list[Finding]
    counts: dict[str, int]

    @property
    def passes(self) -> bool:
        return not self.errors and not self.open_blocking


def check_review(review: Review, spec_path: Path) -> ReviewReport:
    spec = SlideSpec.model_validate_json(spec_path.read_text())
    ids = [c.id for c in claims(spec)]
    errors = []
    if review.spec_sha256 != sha256_file(spec_path):
        errors.append("spec_sha256 doesn't match the spec: the spec changed after the review")
    if review.source_sha256 != spec.pub.source_sha256:
        errors.append("source_sha256 doesn't match the spec")
    missing = [i for i in ids if i not in review.verdicts]
    unknown = [i for i in review.verdicts if i not in ids]
    if missing:
        errors.append(f"{len(missing)} claim(s) without a verdict, e.g. {missing[:3]}")
    if unknown:
        errors.append(f"{len(unknown)} verdict(s) for unknown claims, e.g. {unknown[:3]}")
    bad_refs = [f.claim for f in review.findings if f.claim not in ids]
    if bad_refs:
        errors.append(f"finding(s) for unknown claims: {bad_refs[:3]}")
    for claim, v in review.verdicts.items():
        open_f = [f for f in review.findings if f.claim == claim and f.status == "open"]
        if (
            v != "pass"
            and not open_f
            and not any(f.claim == claim and f.status == "accepted" for f in review.findings)
        ):
            errors.append(f"{claim}: verdict {v} but no open or accepted finding explains it")
    counts = {v: sum(1 for x in review.verdicts.values() if x == v) for v in
              ("pass", "minor", "major", "critical")}  # fmt: skip
    blocking = [
        f for f in review.findings if f.status == "open" and f.verdict in ("major", "critical")
    ]
    return ReviewReport(errors, blocking, counts)
