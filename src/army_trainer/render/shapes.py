"""Drawing helpers shared by the pattern renderers (native python-pptx shapes only)."""

from __future__ import annotations

import math
from dataclasses import dataclass, field

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import theme as t

P = t.PALETTE
RECT = MSO_SHAPE.RECTANGLE
CONTENT_X = t.MARGIN
CONTENT_W = t.SLIDE_W - 2 * t.MARGIN
CONTENT_Y = t.CONTENT_TOP
CONTENT_H = t.CONTENT_BOTTOM - t.CONTENT_TOP
_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
CHAR_EM = 0.5  # average Arial glyph width in em; bold adds 0.05 in est_lines


@dataclass
class Para:
    text: str
    size: int = t.BODY_PT
    bold: bool = False
    color: str = P.army_black
    after: int = 0  # space after, pt


@dataclass
class Ctx:
    """Per-render context: deck facts plus any text that could not be made to fit."""

    short_name: str = ""
    disclaimer: str = t.DISCLAIMER
    warnings: list[str] = field(default_factory=list)


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def flat(shp) -> None:
    """No shadow: empty effectLst for PowerPoint, effectRef idx 0 for LibreOffice."""
    shp.shadow.inherit = False
    for ref in shp._element.xpath(".//a:effectRef"):
        ref.set("idx", "0")


def est_lines(text: str, width_in: float, pt: float, bold: bool = False) -> int:
    per_line = max(1, int(width_in * 72 / (pt * (CHAR_EM + (0.05 if bold else 0)))))
    return sum(max(1, math.ceil(len(line) / per_line)) for line in text.split("\n"))


def est_height(paras: list[Para], width_in: float, delta: int = 0) -> float:
    """Estimated text height in inches at ``delta`` points smaller than each paragraph's size."""
    total = 0.0
    for p in paras:
        pt = p.size - delta
        total += est_lines(p.text, width_in, pt, p.bold) * pt * 1.2 / 72 + p.after / 72
    return total


def box(  # noqa: PLR0913
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str | list[Para] = "",
    *,
    ctx: Ctx | None = None,
    size: int = t.BODY_PT,
    bold: bool = False,
    color: str = P.army_black,
    fill: str | None = None,
    line: str | None = None,
    line_pt: float = 1.5,
    align: str = "l",
    anchor: str = "m",
    shape=RECT,
    floor: int = t.MIN_BODY_PT,
    pad: float = 0.1,
):
    """Add a flat shape with text. Text shrinks (never below ``floor`` pt) to fit; a warning is
    recorded on ``ctx`` if it still doesn't."""
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    flat(shp)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(line_pt)
    else:
        shp.line.fill.background()
    paras = [Para(text, size, bold, color)] if isinstance(text, str) else text
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = _ANCHOR[anchor]
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    if not paras or not any(p.text for p in paras):
        return shp
    inner_w, inner_h = w - 2 * pad, h - 0.1
    delta = 0
    while (
        est_height(paras, inner_w, delta) > inner_h and min(p.size for p in paras) - delta > floor
    ):
        delta += 1
    if est_height(paras, inner_w, delta) > inner_h and ctx is not None:
        ctx.warnings.append(f"text may overflow: {paras[0].text[:40]!r}")
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = _ALIGN[align]
        if p.after:
            para.space_after = Pt(p.after)
        run = para.add_run()
        run.text = p.text
        run.font.name = t.FONT
        run.font.size = Pt(max(p.size - delta, 1))
        run.font.bold = p.bold
        run.font.color.rgb = _rgb(p.color)
    return shp


def line(slide, x1, y1, x2, y2, color=P.army_black, pt=2.0):
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = _rgb(color)
    c.line.width = Pt(pt)
    return c


def check_mark(slide, x, y, size, color=P.do_green, pt=4.0):
    """Open polyline check mark inside a ``size`` x ``size`` square at (x, y)."""
    s = lambda v: int(Inches(v * size))  # noqa: E731
    fb = slide.shapes.build_freeform(Inches(x) + s(0.15), Inches(y) + s(0.55))
    fb.add_line_segments(
        [(Inches(x) + s(0.4), Inches(y) + s(0.8)), (Inches(x) + s(0.85), Inches(y) + s(0.2))],
        close=False,
    )
    shp = fb.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = _rgb(color)
    shp.line.width = Pt(pt)
    flat(shp)
    return shp


def cross_mark(slide, x, y, size, color=P.dont_red, pt=4.0):
    a, b = 0.2 * size, 0.8 * size
    line(slide, x + a, y + a, x + b, y + b, color, pt)
    line(slide, x + a, y + b, x + b, y + a, color, pt)


def set_title(slide, text: str) -> None:
    """Fill the layout's title placeholder; shrink long titles so two lines fit the bar."""
    ph = slide.shapes.title
    ph.text_frame.text = text
    pt = 28 if len(text) <= 45 else 24 if len(text) <= 52 else 22
    for r in ph.text_frame.paragraphs[0].runs:
        r.font.size = Pt(pt)


def drop_empty_placeholders(slide) -> None:
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def emu(v: float) -> Emu:
    return Emu(int(Inches(v)))
