"""Text-led patterns: title, at_a_glance, takeaways, whats_new, divider, key_idea, checklist,
big_numbers, comparison, closing."""

from __future__ import annotations

from pptx.util import Inches

from .. import theme as t
from ..shapes import (
    CONTENT_H,
    CONTENT_W,
    CONTENT_X,
    CONTENT_Y,
    Ctx,
    P,
    box,
    check_mark,
    drop_empty_placeholders,
    set_title,
)

GAP = 0.2


def _rows(n: int, max_h: float = 1.15) -> tuple[float, float]:
    """Row height and gap for ``n`` stacked rows filling the content area."""
    h = min(max_h, (CONTENT_H - GAP * (n - 1)) / n)
    return h, GAP


def title(slide, s, ctx: Ctx) -> None:
    slide.shapes.title.text_frame.text = s.title
    sub = slide.placeholders[1]
    sub.text_frame.text = s.subtitle or ""
    sub.top, sub.height = Inches(4.65), Inches(0.7)
    # the disclaimer fills the lower part of the slide: gold bar + panel, body-size text
    box(slide, t.MARGIN, 5.6, 0.18, 1.3, "", fill=P.army_gold)
    box(
        slide, t.MARGIN + 0.18, 5.6, t.SLIDE_W - 2 * t.MARGIN - 0.18, 1.3, ctx.disclaimer,
        ctx=ctx, size=20, bold=True, fill=P.pale_gray, pad=0.3,
    )  # fmt: skip
    drop_empty_placeholders(slide)


def stat_tiles(slide, stats, ctx: Ctx, y: float, h: float, big_pt: int = 54) -> None:
    n = len(stats)
    w = (CONTENT_W - GAP * (n - 1)) / n
    for i, st in enumerate(stats):
        x = CONTENT_X + i * (w + GAP)
        box(slide, x, y, w, h, "", fill=P.army_gold)
        box(
            slide,
            x,
            y + 0.1,
            w,
            h * 0.5,
            st.value,
            ctx=ctx,
            size=big_pt,
            bold=True,
            align="c",
            floor=24,
            one_line=True,
        )
        box(
            slide,
            x + 0.1,
            y + h * 0.55,
            w - 0.2,
            h * 0.4,
            st.label,
            ctx=ctx,
            size=20,
            align="c",
            anchor="t",
        )


def at_a_glance(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    y = CONTENT_Y
    for tag, item in (("PURPOSE", s.purpose), ("APPLIES TO", s.applies_to)):
        box(
            slide,
            CONTENT_X,
            y,
            2.2,
            0.9,
            tag,
            size=18,
            bold=True,
            color=P.white,
            fill=P.army_black,
            align="c",
        )
        box(
            slide,
            CONTENT_X + 2.2,
            y,
            CONTENT_W - 2.2,
            0.9,
            item.text,
            ctx=ctx,
            size=20,
            fill=P.pale_gray,
        )
        y += 1.1
    if s.stats:
        stat_tiles(slide, s.stats, ctx, y + 0.25, CONTENT_Y + CONTENT_H - y - 0.25)


def takeaways(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    h, g = _rows(len(s.items))
    for i, it in enumerate(s.items):
        y = CONTENT_Y + i * (h + g)
        box(
            slide,
            CONTENT_X,
            y,
            h,
            h,
            str(i + 1),
            size=36,
            bold=True,
            color=P.army_gold,
            fill=P.army_black,
            align="c",
        )
        box(
            slide,
            CONTENT_X + h,
            y,
            CONTENT_W - h,
            h,
            it.text,
            ctx=ctx,
            size=22,
            fill=P.pale_gray,
            pad=0.25,
        )


def whats_new(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    h, g = _rows(len(s.items), 1.0)
    for i, it in enumerate(s.items):
        y = CONTENT_Y + i * (h + g)
        box(slide, CONTENT_X, y, 0.18, h, "", fill=P.army_gold)
        box(
            slide,
            CONTENT_X + 0.18,
            y,
            CONTENT_W - 0.18,
            h,
            it.text,
            ctx=ctx,
            size=22,
            fill=P.pale_gray,
            pad=0.25,
        )


def divider(slide, s, ctx: Ctx) -> None:
    slide.shapes.title.text_frame.text = s.title
    body = slide.placeholders[1]
    body.text_frame.text = s.blurb.text if s.blurb else ""
    if s.number:
        box(
            slide,
            t.MARGIN + 0.3,
            0.5,
            4.0,
            1.4,
            s.number,
            size=88,
            bold=True,
            color=P.army_gold,
            floor=48,
        )
    drop_empty_placeholders(slide)


def key_idea(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    pts = s.points
    hero_h = CONTENT_H if not pts else 2.9
    box(slide, CONTENT_X, CONTENT_Y, 0.2, hero_h, "", fill=P.army_gold)
    box(
        slide,
        CONTENT_X + 0.2,
        CONTENT_Y,
        CONTENT_W - 0.2,
        hero_h,
        s.statement.text,
        ctx=ctx,
        size=32,
        bold=True,
        fill=P.pale_gray,
        pad=0.4,
    )
    if pts:
        n = len(pts)
        w = (CONTENT_W - GAP * (n - 1)) / n
        y = CONTENT_Y + hero_h + 0.3
        for i, it in enumerate(pts):
            box(
                slide,
                CONTENT_X + i * (w + GAP),
                y,
                w,
                CONTENT_Y + CONTENT_H - y,
                it.text,
                ctx=ctx,
                size=20,
                line=P.olive_gray,
                anchor="t",
                pad=0.2,
            )


def checklist(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    h, g = _rows(len(s.items), 0.85)
    for i, it in enumerate(s.items):
        y = CONTENT_Y + i * (h + g)
        box(slide, CONTENT_X, y, h, h, "", line=P.do_green, line_pt=2.5)
        check_mark(slide, CONTENT_X, y, h)
        box(slide, CONTENT_X + h + 0.2, y, CONTENT_W - h - 0.2, h, it.text, ctx=ctx, size=22)


def big_numbers(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    caveat = getattr(s, "caveat", None)
    if caveat:
        # one cited line under the tiles, at body size (not footer size)
        stat_tiles(slide, s.stats, ctx, CONTENT_Y + 0.1, 3.9, big_pt=72)
        y = CONTENT_Y + 4.2
        box(
            slide,
            CONTENT_X,
            y,
            CONTENT_W,
            CONTENT_Y + CONTENT_H - y,
            caveat.text,
            ctx=ctx,
            size=20,
            fill=P.pale_gray,
            pad=0.25,
        )
    else:
        stat_tiles(slide, s.stats, ctx, CONTENT_Y + 0.4, 4.2, big_pt=72)


def comparison(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    n = len(s.columns)
    w = (CONTENT_W - GAP * (n - 1)) / n
    fills = [P.army_black, P.olive_gray, P.mid_gray]
    for i, col in enumerate(s.columns):
        x = CONTENT_X + i * (w + GAP)
        box(
            slide,
            x,
            CONTENT_Y,
            w,
            0.9,
            col.heading,
            ctx=ctx,
            size=22,
            bold=True,
            color=P.white,
            fill=fills[i],
            align="c",
        )
        body_y = CONTENT_Y + 0.9
        k = len(col.points)
        ph = min(2.2, (CONTENT_H - 0.9 - GAP * k) / k)  # short columns: don't stretch panels
        for j, it in enumerate(col.points):
            box(
                slide,
                x,
                body_y + GAP + j * (ph + GAP),
                w,
                ph,
                it.text,
                ctx=ctx,
                size=24,
                fill=P.pale_gray,
                pad=0.2,
            )


def closing(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    h, g = _rows(len(s.items), 1.0)
    for i, it in enumerate(s.items):
        y = CONTENT_Y + i * (h + g)
        box(slide, CONTENT_X, y, 0.18, h, "", fill=P.army_gold)
        box(
            slide,
            CONTENT_X + 0.18,
            y,
            CONTENT_W - 0.18,
            h,
            it.text,
            ctx=ctx,
            size=22,
            fill=P.pale_gray,
            pad=0.25,
        )
    y = CONTENT_Y + len(s.items) * (h + g) + 0.1
    box(
        slide,
        CONTENT_X,
        y,
        CONTENT_W,
        min(1.2, CONTENT_Y + CONTENT_H - y),
        ctx.disclaimer,
        ctx=ctx,
        size=18,
        bold=True,
        fill=P.light_gray,
        pad=0.25,
    )
