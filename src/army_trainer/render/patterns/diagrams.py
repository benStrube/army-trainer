"""Diagram-led patterns: process_flow, cycle, roles, timeline, do_dont, table, decision_tree,
key_terms, acronyms."""

from __future__ import annotations

import math

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .. import theme as t
from ..shapes import (
    CONTENT_H,
    CONTENT_W,
    CONTENT_X,
    CONTENT_Y,
    RECT,
    Ctx,
    P,
    Para,
    _rgb,
    arc_arrow,
    box,
    check_mark,
    cross_mark,
    elbow,
    est_lines,
    line,
    set_title,
)

GAP = 0.2
BOTTOM = CONTENT_Y + CONTENT_H


def process_flow(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    n = len(s.steps)
    if n <= 4:  # horizontal chevrons, detail underneath
        w = CONTENT_W / n
        for i, st in enumerate(s.steps):
            x = CONTENT_X + i * w
            box(
                slide,
                x,
                CONTENT_Y + 0.3,
                w + (0.15 if i < n - 1 else 0),
                1.6,
                st.label,
                ctx=ctx,
                size=22,
                bold=True,
                color=P.white,
                fill=P.army_black,
                align="c",
                shape=MSO_SHAPE.PENTAGON if i == n - 1 else MSO_SHAPE.CHEVRON,
                pad=0.05,
                floor=16,
            )
            box(
                slide,
                x + 0.1,
                CONTENT_Y + 0.3 + 1.6,
                1.0,
                0.5,
                str(i + 1),
                size=24,
                bold=True,
                color=P.gold_dark,
                anchor="m",
                pad=0,
            )
            if st.detail:
                box(
                    slide,
                    x + 0.1,
                    CONTENT_Y + 2.5,
                    w - 0.3,
                    2.6,
                    st.detail,
                    ctx=ctx,
                    size=20,
                    anchor="t",
                    pad=0.05,
                )
    else:  # numbered rows with a spine
        h = (CONTENT_H - GAP * (n - 1)) / n
        line(
            slide,
            CONTENT_X + h / 2,
            CONTENT_Y + h / 2,
            CONTENT_X + h / 2,
            CONTENT_Y + CONTENT_H - h / 2,
            P.light_gray,
            6,
        )
        for i, st in enumerate(s.steps):
            y = CONTENT_Y + i * (h + GAP)
            box(
                slide,
                CONTENT_X,
                y,
                h,
                h,
                str(i + 1),
                size=24,
                bold=True,
                fill=P.army_gold,
                align="c",
                shape=MSO_SHAPE.OVAL,
                pad=0,
            )
            paras = [Para(st.label, 22, True)]
            if st.detail:
                paras.append(Para(st.detail, 18))
            box(slide, CONTENT_X + h + 0.25, y, CONTENT_W - h - 0.25, h, paras, ctx=ctx, floor=16)


def cycle(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    n = len(s.steps)
    d = 3.8
    cx, cy = CONTENT_X + 0.4 + d / 2, CONTENT_Y + CONTENT_H / 2
    nd = 0.8
    gap = math.asin((nd / 2 + 0.12) / (d / 2))  # leave room for the numbered nodes
    for i in range(n):
        a0 = -math.pi / 2 + 2 * math.pi * i / n + gap
        a1 = -math.pi / 2 + 2 * math.pi * (i + 1) / n - gap
        arc_arrow(slide, cx, cy, d / 2, a0, a1)
    if s.center_label:
        box(
            slide,
            cx - 1.3,
            cy - 0.7,
            2.6,
            1.4,
            s.center_label,
            ctx=ctx,
            size=22,
            bold=True,
            align="c",
        )
    for i in range(n):
        a = -math.pi / 2 + 2 * math.pi * i / n
        box(
            slide,
            cx + d / 2 * math.cos(a) - nd / 2,
            cy + d / 2 * math.sin(a) - nd / 2,
            nd,
            nd,
            str(i + 1),
            size=24,
            bold=True,
            fill=P.army_black,
            color=P.army_gold,
            align="c",
            shape=MSO_SHAPE.OVAL,
            pad=0,
        )
    lx = cx + d / 2 + 0.7
    h = (CONTENT_H - GAP * (n - 1)) / n
    for i, st in enumerate(s.steps):
        paras = [Para(f"{i + 1}. {st.label}", 22, True)]
        if st.detail:
            paras.append(Para(st.detail, 18))
        box(
            slide,
            lx,
            CONTENT_Y + i * (h + GAP),
            CONTENT_X + CONTENT_W - lx,
            h,
            paras,
            ctx=ctx,
            fill=P.pale_gray,
            pad=0.2,
        )


def _role_paras(r) -> list[Para]:
    name = f"{r.name} ({r.abbreviation})" if r.abbreviation else r.name
    return [Para(name, 20, True, after=4)] + [Para("• " + d.text, 18) for d in r.duties]


def roles(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    if s.layout == "cards":
        n = len(s.roles)
        cols = 2 if n <= 4 else 3
        rows = math.ceil(n / cols)
        w = (CONTENT_W - GAP * (cols - 1)) / cols
        h = (CONTENT_H - GAP * (rows - 1)) / rows
        for i, r in enumerate(s.roles):
            x, y = CONTENT_X + (i % cols) * (w + GAP), CONTENT_Y + (i // cols) * (h + GAP)
            name = f"{r.name} ({r.abbreviation})" if r.abbreviation else r.name
            box(
                slide,
                x,
                y,
                w,
                0.65,
                name,
                ctx=ctx,
                size=20,
                bold=True,
                color=P.white,
                fill=P.army_black,
                floor=16,
            )
            box(
                slide,
                x,
                y + 0.65,
                w,
                h - 0.65,
                [Para("• " + d.text, 18, after=3) for d in r.duties],
                ctx=ctx,
                fill=P.pale_gray,
                anchor="t",
                floor=14,
            )
        return
    by_key = {r.key: r for r in s.roles}
    depth: dict[str, int] = {}

    def d_of(k: str) -> int:
        if k not in depth:
            p = by_key[k].parent
            depth[k] = 0 if p is None or p not in by_key else d_of(p) + 1
        return depth[k]

    levels: dict[int, list] = {}
    for r in s.roles:
        levels.setdefault(d_of(r.key), []).append(r)
    nl = len(levels)
    bh = (CONTENT_H - 0.5 * (nl - 1)) / nl
    pos: dict[str, tuple[float, float, float]] = {}
    for lv in sorted(levels):
        row = levels[lv]
        w = min(4.0, (CONTENT_W - GAP * (len(row) - 1)) / len(row))
        x0 = CONTENT_X + (CONTENT_W - (w * len(row) + GAP * (len(row) - 1))) / 2
        y = CONTENT_Y + lv * (bh + 0.5)
        for i, r in enumerate(row):
            x = x0 + i * (w + GAP)
            pos[r.key] = (x + w / 2, y, bh)
            if r.parent in pos:
                px, py, ph = pos[r.parent]
                line(slide, px, py + ph, x + w / 2, y, P.mid_gray, 2)
            box(
                slide,
                x,
                y,
                w,
                bh,
                _role_paras(r),
                ctx=ctx,
                fill=P.pale_gray,
                line=P.army_black,
                anchor="t",
                floor=14,
            )


def timeline(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    n = len(s.events)
    ay = CONTENT_Y + CONTENT_H / 2
    line(slide, CONTENT_X, ay, CONTENT_X + CONTENT_W, ay, P.army_black, 5)
    step = CONTENT_W / n
    bw = min(3.2, 2 * step - 0.2)
    bh = CONTENT_H / 2 - 0.55
    for i, e in enumerate(s.events):
        mx = CONTENT_X + step * (i + 0.5)
        up = i % 2 == 0
        box(
            slide,
            mx - 0.16,
            ay - 0.16,
            0.32,
            0.32,
            "",
            fill=P.army_gold,
            line=P.army_black,
            shape=MSO_SHAPE.OVAL,
        )
        line(
            slide,
            mx,
            ay - 0.16 if up else ay + 0.16,
            mx,
            ay - 0.45 if up else ay + 0.45,
            P.army_black,
            2,
        )
        bx = min(max(mx - bw / 2, CONTENT_X), CONTENT_X + CONTENT_W - bw)
        by = ay - 0.45 - bh if up else ay + 0.45
        box(
            slide,
            bx,
            by,
            bw,
            0.55,
            e.when,
            ctx=ctx,
            size=20,
            bold=True,
            fill=P.army_gold,
            align="c",
            anchor="m",
            pad=0.05,
            floor=16,
        )
        paras = [Para(e.label, 18, True)] + ([Para(e.detail, 16)] if e.detail else [])
        box(
            slide,
            bx,
            by + 0.55,
            bw,
            bh - 0.55,
            paras,
            ctx=ctx,
            fill=P.pale_gray,
            anchor="t",
            floor=14,
        )


def do_dont(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    w = (CONTENT_W - GAP) / 2
    for col, (head, color, items, mark) in enumerate(
        [("DO", P.do_green, s.do, check_mark), ("DON'T", P.dont_red, s.dont, cross_mark)]
    ):
        x = CONTENT_X + col * (w + GAP)
        box(
            slide,
            x,
            CONTENT_Y,
            w,
            0.7,
            head,
            size=24,
            bold=True,
            color=P.white,
            fill=color,
            align="c",
        )
        n = len(items)
        h = min(1.3, (CONTENT_H - 0.7 - GAP * n) / n)
        for i, it in enumerate(items):
            y = CONTENT_Y + 0.7 + GAP + i * (h + GAP)
            box(slide, x, y, w, h, "", fill=P.pale_gray)
            mark(slide, x + 0.1, y + (h - 0.6) / 2, 0.6, color, 4)
            box(slide, x + 0.8, y, w - 0.9, h, it.text, ctx=ctx, size=20, anchor="m")


def table(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    nr, nc = len(s.rows) + 1, len(s.columns)
    weights = [
        max([len(s.columns[c])] + [len(r.cells[c]) for r in s.rows if c < len(r.cells)]) + 6
        for c in range(nc)
    ]
    widths = [CONTENT_W * wt / sum(weights) for wt in weights]
    pt = 20
    while pt > 14:
        total = 0.6 + sum(
            max(est_lines(r.cells[c], widths[c] - 0.2, pt) for c in range(nc)) * pt * 1.2 / 72 + 0.2
            for r in s.rows
        )
        if total <= CONTENT_H:
            break
        pt -= 1
    gf = slide.shapes.add_table(
        nr, nc, Inches(CONTENT_X), Inches(CONTENT_Y), Inches(CONTENT_W), Inches(0.6)
    )
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for c in range(nc):
        tbl.columns[c].width = Inches(widths[c])
    tbl.rows[0].height = Inches(0.6)

    def fill_cell(cell, text, bold, color, bg):
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(bg)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.1)
        tf = cell.text_frame
        tf.word_wrap = True
        tf.text = text
        for para in tf.paragraphs:
            para.alignment = PP_ALIGN.LEFT
            for r in para.runs:
                r.font.name, r.font.size, r.font.bold = t.FONT, Pt(pt), bold
                r.font.color.rgb = _rgb(color)

    for c, name in enumerate(s.columns):
        fill_cell(tbl.cell(0, c), name, True, P.white, P.army_black)
    for r_i, row in enumerate(s.rows, start=1):
        lines = max(est_lines(row.cells[c], widths[c] - 0.2, pt) for c in range(nc))
        tbl.rows[r_i].height = Inches(lines * pt * 1.2 / 72 + 0.2)
        for c in range(nc):
            fill_cell(
                tbl.cell(r_i, c),
                row.cells[c] if c < len(row.cells) else "",
                c == 0,
                P.army_black,
                P.pale_gray if r_i % 2 else P.white,
            )
    total = sum(r.height for r in tbl.rows) / 914400
    if total > CONTENT_H:
        ctx.warnings.append(f"table may overflow slide {s.id} ({total:.1f} in)")


def decision_tree(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    nodes = {n.key: n for n in s.nodes}
    xs: dict[str, float] = {}
    depth: dict[str, int] = {}
    leaf = [0]

    def walk(k: str, d: int) -> float:
        depth[k] = d
        kids = [c for c in (nodes[k].yes, nodes[k].no) if c in nodes]
        if not kids:
            xs[k] = float(leaf[0])
            leaf[0] += 1
        else:
            xs[k] = sum(walk(c, d + 1) for c in kids) / len(kids)
        return xs[k]

    walk(s.root, 0)
    levels = max(depth.values()) + 1
    nleaf = max(leaf[0], 1)
    cw = CONTENT_W / nleaf
    bw = min(3.6, cw - 0.15)
    gap_y = 0.75
    bh = (CONTENT_H - gap_y * (levels - 1)) / levels

    def at(k: str) -> tuple[float, float]:
        return CONTENT_X + (xs[k] + 0.5) * cw - bw / 2, CONTENT_Y + depth[k] * (bh + gap_y)

    for k, n in nodes.items():
        if k not in xs:
            continue
        x, y = at(k)
        for label, child in (("YES", n.yes), ("NO", n.no)):
            if child in nodes:
                cx, cy = at(child)
                ym = elbow(slide, x + bw / 2, y + bh, cx + bw / 2, cy, P.army_black, 3)
                color = P.do_green if label == "YES" else P.dont_red
                # the answer sits on the horizontal run, next to the child it leads to
                px = cx + bw / 2 + (-0.45 if cx > x else 0.45 if cx < x else 0.0) * 1.0
                box(
                    slide, px - 0.35, ym - 0.2, 0.7, 0.4, label, size=14, bold=True,
                    color=P.white, fill=color, align="c", pad=0.02, floor=12,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                )  # fmt: skip
    for k, n in nodes.items():
        if k not in xs:
            continue
        x, y = at(k)
        q = n.kind == "question"
        box(
            slide,
            x,
            y,
            bw,
            bh,
            n.text,
            ctx=ctx,
            size=18,
            bold=q,
            color=P.white if q else P.army_black,
            fill=P.army_black if q else P.army_gold,
            align="c",
            floor=16,
            pad=0.1,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE if q else RECT,
        )


def key_terms(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    n = len(s.terms)
    cols = 1 if n <= 3 else 2
    rows = math.ceil(n / cols)
    w = (CONTENT_W - GAP * (cols - 1)) / cols
    h = (CONTENT_H - GAP * (rows - 1)) / rows
    head = 0.5
    for i, tm in enumerate(s.terms):
        x, y = CONTENT_X + (i % cols) * (w + GAP), CONTENT_Y + (i // cols) * (h + GAP)
        # card: the term on a black tab, the plain-language definition below it
        box(slide, x, y, w, head, tm.term, ctx=ctx, size=22, bold=True, color=P.army_gold,
            fill=P.army_black, pad=0.2, floor=18)  # fmt: skip
        box(slide, x, y + head, w, h - head, tm.definition, ctx=ctx, size=20, fill=P.pale_gray,
            anchor="t", pad=0.2, floor=16)  # fmt: skip


def acronyms(slide, s, ctx: Ctx) -> None:
    set_title(slide, s.title)
    half = math.ceil(len(s.entries) / 2)
    cols = [s.entries[:half], s.entries[half:]]
    w = (CONTENT_W - GAP) / 2
    aw = 1.5
    pt = 18
    while pt > 12:
        need = max(
            sum(max(1, est_lines(e.meaning, w - aw - 0.1, pt)) * pt * 1.2 / 72 + 0.12 for e in col)
            for col in cols
        )
        if need <= CONTENT_H:
            break
        pt -= 1
    for ci, col in enumerate(cols):
        x, y = CONTENT_X + ci * (w + GAP), CONTENT_Y
        for i, e in enumerate(col):
            h = max(1, est_lines(e.meaning, w - aw - 0.1, pt)) * pt * 1.2 / 72 + 0.12
            if i % 2 == 0:
                box(slide, x, y, w, h, "", fill=P.pale_gray)
            box(slide, x, y, aw, h, e.abbreviation, size=pt, bold=True, floor=12, pad=0.08, ctx=ctx)
            box(slide, x + aw, y, w - aw, h, e.meaning, size=pt, floor=12, pad=0.05, ctx=ctx)
            y += h
        if y > BOTTOM + 0.05:
            ctx.warnings.append("acronym list may overflow")
