"""Render slide-spec slides onto the base template (WP 3.2: patterns only; 3.3 adds footers,
notes and the full-deck assembly)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from ..plan.spec import SlideSpec
from .patterns import REGISTRY
from .shapes import Ctx
from .template import DEFAULT_PATH


def render_slides(
    spec: SlideSpec, out: Path, template: Path = DEFAULT_PATH, only: set[str] | None = None
) -> Ctx:
    """Draw every slide (or those whose id is in ``only``). Returns the context with any
    text-fit warnings."""
    prs = Presentation(template)
    layouts = {la.name: la for la in prs.slide_layouts}
    ctx = Ctx(short_name=spec.pub.short_name, disclaimer=spec.disclaimer)
    for s in spec.slides:
        if only and s.id not in only:
            continue
        layout_name, fn = REGISTRY[s.pattern]
        fn(prs.slides.add_slide(layouts[layout_name]), s, ctx)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return ctx
