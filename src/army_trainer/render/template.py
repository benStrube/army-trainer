"""Build ``templates/base.pptx``: a 16:9 master with the D9 look.

Layouts kept: ``Title Slide``, ``Section Header``, ``Content`` (title bar + gold rule, a free body
area for the pattern renderers), ``Blank``. The footer and slide-number placeholders exist on each
layout; the deck assembler (WP 3.3) copies them onto slides.
"""

from __future__ import annotations

import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from . import theme as t

DEFAULT_PATH = Path(__file__).resolve().parents[3] / "templates" / "base.pptx"
KEEP = {
    "Title Slide": "Title Slide",
    "Section Header": "Section Header",
    "Title Only": "Content",
    "Blank": "Blank",
}


def _rect(layout, x, y, w, h, color, name):
    """Add a flat rectangle behind the placeholders (LayoutShapes has no add_shape)."""
    tree = layout.shapes._spTree
    sid = max([int(i) for i in tree.xpath("//p:cNvPr/@id")] + [1]) + 1
    e = lambda v: int(Inches(v))  # noqa: E731
    xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/>'
        '<p:nvPr userDrawn="1"/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>'
        "</p:sp>"
    )
    shape = etree.fromstring(xml)
    first_ph = next((sp for sp in tree.findall(qn("p:sp")) if sp.xpath(".//p:ph")), None)
    if first_ph is not None:
        first_ph.addprevious(shape)  # draw order: earlier rects first, placeholders on top
    else:
        tree.append(shape)


def _place(
    ph,
    x,
    y,
    w,
    h,
    size,
    color,
    bold=False,
    caps=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.MIDDLE,
):
    ph.left, ph.top, ph.width, ph.height = (Inches(v) for v in (x, y, w, h))
    tf = ph.text_frame
    tf.vertical_anchor = anchor
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.alignment = align
    # layout-level defaults so slides inherit them
    lst = tf._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lstStyle")
    for child in list(lst):
        lst.remove(child)
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    lvl = etree.SubElement(
        lst, f"{{{ns}}}lvl1pPr", algn="l" if align == PP_ALIGN.LEFT else "r", marL="0", indent="0"
    )
    etree.SubElement(lvl, f"{{{ns}}}buNone")
    attrs = {"sz": str(int(size * 100)), "b": "1" if bold else "0"}
    if caps:
        attrs["cap"] = "all"
        attrs["spc"] = "100"
    d = etree.SubElement(lvl, f"{{{ns}}}defRPr", **attrs)
    fill = etree.SubElement(d, f"{{{ns}}}solidFill")
    etree.SubElement(fill, f"{{{ns}}}srgbClr", val=color)
    etree.SubElement(d, f"{{{ns}}}latin", typeface=t.FONT)


def _style_master_text(prs) -> None:
    """Master title/body styles: Arial, D9 colors, bold caps titles."""
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    root = prs.slide_master._element
    pns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    title = root.find(f".//{{{pns}}}titleStyle/{{{ns}}}lvl1pPr")
    body = root.find(f".//{{{pns}}}bodyStyle")
    for ppr, size, bold, caps in [(title, t.TITLE_PT, True, True)]:
        d = ppr.find(f"{{{ns}}}defRPr")
        d.set("sz", str(size * 100))
        d.set("b", "1" if bold else "0")
        if caps:
            d.set("cap", "all")
            d.set("spc", "100")
        for c in list(d):
            d.remove(c)
            f = etree.SubElement(d, f"{{{ns}}}solidFill")
        etree.SubElement(f, f"{{{ns}}}srgbClr", val=t.PALETTE.army_black)
        etree.SubElement(d, f"{{{ns}}}latin", typeface=t.FONT)
    for i, lvl in enumerate(body):
        d = lvl.find(f"{{{ns}}}defRPr")
        if d is None:
            continue
        d.set("sz", str(max(t.MIN_BODY_PT, t.BODY_PT - 2 * i) * 100))
        for c in list(d):
            d.remove(c)
            f = etree.SubElement(d, f"{{{ns}}}solidFill")
        etree.SubElement(f, f"{{{ns}}}srgbClr", val=t.PALETTE.army_black)
        etree.SubElement(d, f"{{{ns}}}latin", typeface=t.FONT)


def _patch_theme(prs) -> None:
    part = prs.slide_master.part.part_related_by(RT.THEME)
    xml = part.blob.decode("utf-8")
    p = t.PALETTE
    colors = {
        "dk1": p.army_black,
        "lt1": p.white,
        "dk2": p.olive_gray,
        "lt2": p.light_gray,
        "accent1": p.army_gold,
        "accent2": p.olive_gray,
        "accent3": p.mid_gray,
        "accent4": p.gold_dark,
        "accent5": p.do_green,
        "accent6": p.dont_red,
        "hlink": p.gold_dark,
        "folHlink": p.mid_gray,
    }
    for name, val in colors.items():
        xml = re.sub(
            rf"(<a:{name}>).*?(</a:{name}>)", rf'\1<a:srgbClr val="{val}"/>\2', xml, flags=re.S
        )
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', rf"\1{t.FONT}\2", xml)
    xml = re.sub(r'<a:clrScheme name="[^"]*"', '<a:clrScheme name="Army trainer (D9)"', xml)
    part._blob = xml.encode("utf-8")


def build_template(path: Path = DEFAULT_PATH) -> Path:
    p = t.PALETTE
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(t.SLIDE_W), Inches(t.SLIDE_H)
    _patch_theme(prs)
    _style_master_text(prs)

    # drop unused layouts
    for layout in list(prs.slide_layouts):
        if layout.name not in KEEP:
            prs.slide_layouts.remove(layout)
    for layout in prs.slide_layouts:
        layout.name = KEEP[layout.name]
        for ph in list(layout.placeholders):
            if ph.placeholder_format.type == 16:  # date: not used
                ph._element.getparent().remove(ph._element)

    W, M = t.SLIDE_W, t.MARGIN
    for layout in prs.slide_layouts:
        by_type = {ph.placeholder_format.type: ph for ph in layout.placeholders}
        name = layout.name
        footer_color = p.mid_gray
        if name == "Content":
            _rect(layout, 0, 0, W, t.TITLE_BAR_H, p.army_black, "Title bar")
            _rect(layout, 0, t.TITLE_BAR_H, W, t.GOLD_RULE_H, p.army_gold, "Gold rule")
            _place(
                by_type[1],
                M,
                0.12,
                W - 2 * M,
                t.TITLE_BAR_H - 0.24,
                t.TITLE_PT,
                p.white,
                bold=True,
                caps=True,
            )
        elif name == "Title Slide":
            _rect(layout, 0, 0, W, 4.9, p.army_black, "Title panel")
            _rect(layout, 0, 4.9, W, 0.12, p.army_gold, "Gold rule")
            _place(
                by_type[3],
                M + 0.3,
                1.2,
                W - 2 * M - 0.6,
                3.2,
                40,
                p.white,
                bold=True,
                caps=True,
                anchor=MSO_ANCHOR.BOTTOM,
            )
            _place(
                by_type[4],
                M + 0.3,
                5.3,
                W - 2 * M - 0.6,
                1.4,
                22,
                p.army_black,
                anchor=MSO_ANCHOR.TOP,
            )
        elif name == "Section Header":
            _rect(layout, 0, 0, W, t.SLIDE_H, p.olive_gray, "Divider background")
            _rect(layout, M + 0.3, 4.1, 3.0, 0.1, p.army_gold, "Gold rule")
            _place(
                by_type[1],
                M + 0.3,
                1.6,
                W - 2 * M - 0.6,
                2.3,
                40,
                p.white,
                bold=True,
                caps=True,
                anchor=MSO_ANCHOR.BOTTOM,
            )
            _place(
                by_type[2], M + 0.3, 4.4, W - 2 * M - 0.6, 1.6, 22, p.white, anchor=MSO_ANCHOR.TOP
            )
            footer_color = p.white
        for ph_type, x, w, align in [
            (15, M, W - 2 * M - 1.5, PP_ALIGN.LEFT),
            (13, W - M - 1.0, 1.0, PP_ALIGN.RIGHT),
        ]:
            if ph_type in by_type:
                _place(
                    by_type[ph_type], x, t.FOOTER_Y, w, 0.3, t.FOOTER_PT, footer_color, align=align
                )
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


if __name__ == "__main__":
    print(build_template())
